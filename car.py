import pandas as pd
from openpyxl import load_workbook
from io import BytesIO
import requests
from sqlalchemy import create_engine
import streamlit as st

# ======================= CONFIGURAÇÃO =======================
SHEET_ID = "1QmCa_NWF9aspVqZSF25ke0luPajyrpei"
XLSX_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

try:
    DB_URL = st.secrets["postgres"]["url"]
except Exception:
    DB_URL = "postgresql://postgres:danielDantas123@db.nqghtobrrvtmstowirix.supabase.co:5432/postgres"

engine = create_engine(DB_URL)
# ============================================================

MARCAS_CONHECIDAS = [
"AUDI", "BMW", "CHEVROLET", "CITROEN", "CITROËN", "FIAT", "FORD",
    "HONDA", "HYUNDAI", "JEEP", "KIA", "MERCEDES", "MERCEDES-BENZ",
    "NISSAN", "PEUGEOT", "RENAULT", "TOYOTA", "VOLKSWAGEN", "VW", "VOLKS", "VOLVO"
]

def limpar_texto(valor):
    if pd.isna(valor) or valor is None:
        return None
    texto = str(valor).strip()
    if texto.lower() in ["", "não informado", "nao informado", "n/i", "nan", "none"]:
        return None
    return texto

def limpar_preco(valor):
    if pd.isna(valor) or valor is None:
        return None

    # Se já veio como número do Excel
    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip()
    texto = texto.replace("R$", "").replace(" ", "")

    # Formato brasileiro: 99.900,00 → 99900.00
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")

    try:
        return float(texto)
    except:
        return None

def limpar_km(valor):
    if pd.isna(valor) or valor is None:
        return None
    try:
        return int(str(valor).replace(".", "").replace(",", "").strip())
    except:
        return None

def main():
    print("Baixando planilha de estoque (.xlsx)...")
    response = requests.get(XLSX_URL)
    response.raise_for_status()

    wb = load_workbook(filename=BytesIO(response.content), data_only=True)

    print("Abas encontradas:", wb.sheetnames)

    if "TABELA DE ESTOQUE" in wb.sheetnames:
        ws = wb["TABELA DE ESTOQUE"]
    else:
        ws = wb.active
        print(f"Usando aba: {ws.title}")

    data = [list(row) for row in ws.iter_rows(values_only=True)]
    df = pd.DataFrame(data)
    print(f"Total de linhas lidas: {len(df)}")

    conn = engine.raw_connection()
    cursor = conn.cursor()
    print("Conectado ao Supabase.")

    marca_atual = None
    inseridos = 0
    atualizados = 0
    pulados = 0

    for idx, row in df.iterrows():
        col0 = limpar_texto(row[0]) if len(row) > 0 else None
        if not col0:
            continue

        # Detecta marca
        if col0.upper() in MARCAS_CONHECIDAS:
            marca_atual = col0.upper()
            print(f"→ Marca: {marca_atual}")
            continue

        # Pula cabeçalhos
        if col0 in ["CARRO", "MODELO", "ESTOQUE MANU AUTOMÓVEIS"]:
            continue

        carro       = limpar_texto(row[0]) if len(row) > 0 else None
        modelo      = limpar_texto(row[1]) if len(row) > 1 else None
        preco       = limpar_preco(row[2]) if len(row) > 2 else None
        ano         = limpar_texto(row[3]) if len(row) > 3 else None
        patio       = limpar_texto(row[4]) if len(row) > 4 else None
        cor         = limpar_texto(row[5]) if len(row) > 5 else None
        combustivel = limpar_texto(row[6]) if len(row) > 6 else None
        placa       = limpar_texto(row[7]) if len(row) > 7 else None
        km          = limpar_km(row[8])    if len(row) > 8 else None
        leilao      = bool(limpar_texto(row[9])) if len(row) > 9 else False

        if not carro:
            pulados += 1
            continue

        # Se não tiver placa, cria uma temporária
        if not placa:
            placa = f"SEM-PLACA-{idx}"

        # Verifica se já existe
        cursor.execute("SELECT id FROM estoque_carros WHERE placa = %s", (placa,))
        existing = cursor.fetchone()

        if existing:
            cursor.execute("""
                UPDATE estoque_carros SET
                    marca = %s,
                    carro = %s,
                    modelo = %s,
                    preco = %s,
                    ano = %s,
                    patio = %s,
                    cor = %s,
                    combustivel = %s,
                    km = %s,
                    leilao = %s,
                    updated_at = NOW()
                WHERE placa = %s
            """, (marca_atual, carro, modelo, preco, ano, patio, cor, combustivel, km, leilao, placa))
            atualizados += 1
        else:
            cursor.execute("""
                INSERT INTO estoque_carros (
                    marca, carro, modelo, preco, ano, patio, cor,
                    combustivel, placa, km, leilao
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """, (marca_atual, carro, modelo, preco, ano, patio, cor, combustivel, placa, km, leilao))
            inseridos += 1

    conn.commit()
    cursor.close()
    conn.close()

    print(f"\n✅ Importação concluída!")
    print(f"Novos carros inseridos: {inseridos}")
    print(f"Carros atualizados:     {atualizados}")
    print(f"Linhas puladas:         {pulados}")

if __name__ == "__main__":
    main()